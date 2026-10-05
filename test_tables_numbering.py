"""Regression tests for tables, numbering, appendices and title page.

Run with: python test_tables_numbering.py (requires the normal converter dependencies).
"""
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

import latex_to_docx as converter


SOURCE = r"""
\documentclass[11pt,a4paper]{article}
\usepackage{tabularx}
\usepackage{longtable}
\usepackage{fancyhdr}
\newcommand{\Programa}{Galapagos Water Program}
\pagestyle{fancy}
\fancyhf{}
\fancyhead[R]{\small\Programa}
\title{%
  \begin{center}
    \includegraphics[height=1.8cm]{logo.png}\\[0.6cm]
    {\Large \textbf{TERMS OF REFERENCE}}\\[0.4cm]
    {\large Subtitle of the assignment}\\[0.4cm]
    {\small Unit name}
  \end{center}}
\author{Company}
\date{October 2026}
\begin{document}
\maketitle

\begin{table}[H]
\centering
\caption{Signature details}
\label{tab:sign}
\begin{tabular}{|p{3cm}|p{3cm}|p{3cm}|}
\hline
\textbf{Prepared by} & \textbf{Position} & \textbf{Signature} \\
\hline
Ana Perez & Specialist & \\
\hline
\end{tabular}
\end{table}

\section{Body}
\label{sec:body}
Body text. See Table~\ref{tab:sign}, Table~\ref{tab:long}, Table~\ref{tab:dates}
and \ref{ann:first}.

Penalty of 3\textperthousand{} per day at 20\textdegree C.

\begin{longtable}{>{\raggedright\arraybackslash}p{2.5cm}>{\raggedright\arraybackslash}p{10cm}}
\caption{Acronyms}\label{tab:long}\\
\textbf{Acronym} & \textbf{Meaning} \\
\hline
\endfirsthead
\textbf{Acronym} & \textbf{Meaning} \\
\hline
\endhead
O\&M & Operation and Maintenance \\
\hline
W\&S & Water and Sanitation \\
\hline
\end{longtable}

\begin{longtable}{p{3cm}p{8cm}}
\toprule
\textbf{Term} & \textbf{Definition} \\
\midrule
\endhead
ASSS & Requirements
spanning two lines \\
\bottomrule
\end{longtable}

\begin{table}[H]
\centering
\caption{Dates}
\label{tab:dates}
\begin{tabularx}{\textwidth}{>{\raggedright\arraybackslash}X>{\raggedright\arraybackslash}p{4cm}>{\centering\arraybackslash}p{1cm}}
\hline
\textbf{Stage} & \textbf{Date} & \textbf{Bar} \\
\hline
Publication & 12 October 2026 & \rule{\linewidth}{1.6ex} \\
\hline
Award & 4 November 2026 & \\
\hline
\end{tabularx}
\end{table}

\begin{longtable}{p{4cm}p{1.1cm}p{1.2cm}p{2cm}}
\caption{Budget}\label{tab:budget}\\
\textbf{Description} & \textbf{Unit} & \textbf{Qty} & \textbf{Total} \\
\hline
\endfirsthead
\textbf{Description} & \textbf{Unit} & \textbf{Qty} & \textbf{Total} \\
\hline
\endhead
\midrule
\multicolumn{4}{r}{\textit{Continúa en la página siguiente}} \\
\endfoot
\bottomrule
\multicolumn{4}{l}{Source: reference prices.} \\
\endlastfoot
\multicolumn{4}{l}{\textbf{DIRECT COSTS}} \\
\hline
Fees & month & 1 & 9,000 \\
\hline
\multicolumn{3}{l}{\textbf{SUBTOTAL FEES}} & \textbf{9,000} \\
\hline
\end{longtable}

\appendix
\renewcommand{\thesection}{Annex \Alph{section}}
\renewcommand{\thesubsection}{\Alph{section}.\arabic{subsection}}

\section{First annex}
\label{ann:first}
Annex text.

\subsection{Annex detail}
Detail text.

\end{document}
"""


def _texts(document):
    return [p.text for p in document.paragraphs]


def _table_with(document, text):
    for table in document.tables:
        for row in table.rows:
            if any(text in cell.text for cell in row.cells):
                return table
    raise AssertionError(f"No table contains {text!r}")


def check_tables_numbering():
    with tempfile.TemporaryDirectory(prefix="latex_converter_tables_") as folder:
        root = Path(folder)
        Image.new("RGB", (100, 30), "white").save(root / "logo.png")
        entry = root / "document.tex"
        output = root / "document.docx"
        entry.write_text(SOURCE, encoding="utf-8")
        command = [sys.executable, str(Path(converter.__file__).resolve()), str(entry), str(output)]
        env = {**os.environ, "PYTHONUTF8": "1", "PYTHONDONTWRITEBYTECODE": "1"}
        subprocess.run(command, check=True, capture_output=True, env=env, timeout=180)
        document = Document(output)
        texts = _texts(document)
        body_text = "\n".join(texts)

        # Title: environment names must not leak and each \\ line is its own paragraph.
        assert not any("center" in t for t in texts), "\\begin{center} leaked into the title"
        title = next(p for p in document.paragraphs if p.text.strip() == "TERMS OF REFERENCE")
        assert title.runs[0].bold
        assert any(t.strip() == "Subtitle of the assignment" for t in texts)
        # article class without the titlepage option: no page break after \maketitle.
        assert not document.element.body.xpath(".//w:br[@w:type='page']"), "unexpected page break"

        # Header macro after a control word (\small\Programa) keeps its first word.
        header = "".join(document.sections[0].header._element.xpath(".//w:t/text()"))
        assert "Galapagos Water Program" in header, header

        # Header row and first data row stay separate when a rule divides them.
        sign = _table_with(document, "Ana Perez")
        assert len(sign.rows) == 2, [c.text for r in sign.rows for c in r.cells]
        assert sign.cell(0, 0).text == "Prepared by"
        assert sign.cell(1, 0).text == "Ana Perez"

        # Escaped ampersands do not split cells.
        acr = _table_with(document, "Operation and Maintenance")
        assert len(acr.columns) == 2
        assert any(row.cells[0].text == "O&M" for row in acr.rows)
        assert any(row.cells[0].text == "W&S" for row in acr.rows)
        assert acr.cell(0, 0).text == "Acronym"

        # A lone \endhead is not cell text; source line breaks inside a cell are spaces.
        terms = _table_with(document, "ASSS")
        assert terms.cell(1, 0).text == "ASSS", repr(terms.cell(1, 0).text)
        assert terms.cell(1, 1).text == "Requirements spanning two lines", repr(terms.cell(1, 1).text)

        # tabularx: width argument is not the column spec; >{...}p{} cells keep leading numbers.
        dates = _table_with(document, "Publication")
        cells = [c.text for r in dates.rows for c in r.cells]
        assert not any(">" in c for c in cells), cells
        assert dates.cell(1, 1).text == "12 October 2026", cells
        assert dates.cell(2, 1).text == "4 November 2026", cells
        # \rule in a cell becomes a shaded cell (e.g. Gantt bars).
        shd = dates.cell(1, 2)._tc.xpath(".//w:shd")
        assert shd and shd[0].get(qn("w:fill")) not in (None, "auto", "FFFFFF"), "rule cell not shaded"

        # \multicolumn spans become merged cells.
        budget = _table_with(document, "SUBTOTAL FEES")
        row = next(r for r in budget.rows if "SUBTOTAL FEES" in r.cells[0].text)
        first_tc = row._tr.tc_lst[0]
        span = first_tc.tcPr.find(qn("w:gridSpan")) if first_tc.tcPr is not None else None
        assert span is not None and span.get(qn("w:val")) == "3"
        assert budget.cell(0, 0).text == "Description"
        assert "DIRECT COSTS" not in budget.cell(0, 0).text
        # longtable foot (\endfoot) only exists at page breaks; the last foot ends the table.
        budget_text = [c.text for r in budget.rows for c in r.cells]
        assert not any("Continúa" in t for t in budget_text), budget_text
        assert budget.rows[-1].cells[0].text == "Source: reference prices.", budget_text
        assert budget.rows[1].cells[0].text == "DIRECT COSTS", budget_text

        # LaTeX never splits a row across pages; a longtable \endhead repeats.
        for table in document.tables:
            for tr in table._tbl.tr_lst:
                assert tr.trPr is not None and tr.trPr.find(qn("w:cantSplit")) is not None
        assert acr._tbl.tr_lst[0].trPr.find(qn("w:tblHeader")) is not None
        assert acr.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next

        # textcomp symbols.
        assert "Penalty of 3‰ per day at 20°C." in body_text, body_text

        # Captions stay with their tables.
        for p in document.paragraphs:
            if p.text.startswith("Table "):
                assert p.paragraph_format.keep_with_next, p.text

        # Table numbers follow source order across table and longtable.
        assert "Table 1: Signature details" in texts
        assert "Table 2: Acronyms" in texts
        assert "Table 3: Dates" in texts
        assert "See Table 1, Table 2, Table 3" in body_text.replace(" ", " ")

        # \appendix + \renewcommand{\thesection}: annex numbering and references.
        assert any(t.startswith("Annex A") and "First annex" in t for t in texts), texts
        assert any(t.startswith("A.1") and "Annex detail" in t for t in texts), texts
        assert "and Annex A." in body_text

    # Column widths include LaTeX's \tabcolsep padding (2 x 6pt) and scale to fit.
    pad_mm = 2 * 6 / 72.27 * 25.4
    widths = converter._parse_col_widths_dxa("p{30mm}p{35mm}", 2)
    assert all(abs(w / 1440 * 25.4 - (t + pad_mm)) < 0.2 for w, t in zip(widths, (30, 35))), widths
    wide = converter._parse_col_widths_dxa("p{10cm}p{10cm}p{1cm}", 3, 9026)
    assert sum(wide) <= 9026 and wide[0] > 5 * wide[2], wide
    bs = chr(92)
    # Thin rules are signature/fill-in lines, thick rules are bars.
    assert converter._rule_fill(bs + "rule{0.2" + bs + "textwidth}{0.4pt}") is None
    assert converter._rule_fill(bs + "rule{" + bs + "linewidth}{1.6ex}") == "000000"
    line = converter._clean("Date: " + bs + "rule{3cm}{0.4pt}")
    assert line.startswith("Date: ") and set(line[6:]) == {"_"} and 10 <= len(line[6:]) <= 20, line

    # A local \setlength{\tabcolsep}{3pt} right before the table is honoured.
    before = "Text." + chr(10) + bs + "begingroup" + chr(10) + bs + "small" + chr(10) + bs + "setlength{" + bs + "tabcolsep}{3pt}" + chr(10)
    assert abs(converter._local_tabcolsep_pt(before, "") - 3) < 0.01
    assert converter._local_tabcolsep_pt("Text." + chr(10) + chr(10) + "More text." + chr(10), "") is None
    assert abs(converter._local_tabcolsep_pt("", bs + "centering " + bs + "setlength{" + bs + "tabcolsep}{4pt}") - 4) < 0.01

    # A table that ends the document must not push an empty page.
    tail = Document()
    tail.add_paragraph("Text")
    tail.add_table(rows=1, cols=1)
    converter._remove_empty_paragraphs_after_tables(tail)
    last = tail.paragraphs[-1]
    assert last._p.getprevious().tag == qn("w:tbl")
    assert last.paragraph_format.space_after == 0 and last.paragraph_format.line_spacing.pt <= 1

    # Escaped ampersands and nested braces are respected when splitting cells.
    assert converter._split_table_cells(r"O\&M & {a & b} & c") == ["O\\&M", "{a & b}", "c"]
    print("OK: tables, numbering, appendices, title and header macros.")


if __name__ == "__main__":
    check_tables_numbering()
