"""Run with: python test_conversion.py (requires the normal converter dependencies)."""
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

import latex_to_docx as converter


def check_conversion():
    source = r"""
\documentclass[12pt,letterpaper]{article}
\usepackage[utf8]{inputenc}
\usepackage[spanish]{babel}
\usepackage{helvet}
\renewcommand{\familydefault}{\sfdefault}
\setlength{\parskip}{7pt}
\fancyhead[L]{\small\textbf{BORRADOR}\\Para revisión}
\fancyhead[R]{\includegraphics[width=37mm]{logo.png}}
\fancyfoot[C]{\thepage}
\AddToShipoutPictureBG{\AtPageLowerLeft{\includegraphics[width=\paperwidth,height=22mm]{logo.png}}}
\newcommand{\pregunta}[2]{\par\begin{samepage}\textbf{Pregunta #1. Consulta:} #2\par\end{samepage}\nopagebreak[3]}
\newcommand{\respuesta}{\textbf{Respuesta:} }
\newcommand{\alias}{\respuesta}
\newcommand{\opcional}[2][Base]{\textbf{#1:} #2}
\newcommand{\cadena}{\respuesta}
\begin{document}
\fontsize{10.5}{12.6}\selectfont
\begin{flushright}\begin{minipage}{63mm}\itshape Dirección\\Ciudad\end{minipage}\end{flushright}
\noindent\textbf{Asunto:} Prueba.

\pregunta{1}{Texto con \textit{formato} y \textbf{llaves {anidadas}}.}
\alias Contenido conservado con precisión $\pm1$\,\%.

\opcional{Texto predeterminado.}

\opcional[Otro]{Texto opcional.}

\begin{center}
\begin{tabular}{l>{\centering\arraybackslash}p{30mm}}
\toprule Hito & Valor\\\midrule P3 & 28\,\%\\\bottomrule
\end{tabular}
\end{center}
\pregunta{2}{Consulta pendiente.}
\respuesta
\end{document}
"""
    with tempfile.TemporaryDirectory(prefix="latex_converter_test_") as folder:
        root = Path(folder)
        Image.new("RGB", (100, 30), "white").save(root / "logo.png")
        entry = root / "document.tex"
        output = root / "document.docx"
        entry.write_text(source, encoding="utf-8")
        command = [sys.executable, str(Path(converter.__file__).resolve()), str(entry), str(output)]
        env = {**os.environ, "PYTHONUTF8": "1", "PYTHONDONTWRITEBYTECODE": "1"}
        subprocess.run(command, check=True, capture_output=True, env=env, timeout=90)
        document = Document(output)
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Pregunta 1. Consulta:" in text and "Pregunta 2. Consulta:" in text
        assert "Contenido conservado con precisión" in text and "llaves anidadas" in text
        assert "Base: Texto predeterminado." in text and "Otro: Texto opcional." in text
        assert "10.512.6" not in text and "TABLE OF CONTENTS" not in text
        assert document.paragraphs[-1].text.strip() == "Respuesta:"
        assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
        for paragraph in document.paragraphs:
            if paragraph.text.startswith(("Pregunta", "Respuesta")):
                assert paragraph.runs[0].bold
                assert paragraph.runs[0].font.name == "Arial"
                assert paragraph.runs[0].font.size.pt == 10.5
        question = next(p for p in document.paragraphs if p.text.startswith("Pregunta 1."))
        assert question.paragraph_format.keep_together and question.paragraph_format.keep_with_next
        assert document.paragraphs[0].paragraph_format.line_spacing.pt == 12.6
        assert document.tables[0].cell(1, 1).text == "28\u202f%"
        # p{30mm} is the text width; Word's column also holds 2 x \tabcolsep (6pt).
        pad_mm = 2 * 6 * 20 / 1440 * 25.4
        assert abs(document.tables[0].columns[1].width.mm - (30 + pad_mm)) < 0.1
        assert document.tables[0].cell(0, 0).paragraphs[0].paragraph_format.keep_with_next
        assert not document.tables[0].cell(1, 0).paragraphs[0].paragraph_format.keep_with_next
        assert document.element.body.xpath(".//m:oMath")
        assert "BORRADOR" in "".join(document.sections[0].header._element.xpath(".//w:t/text()"))
        assert document.sections[0].header._element.xpath(".//pic:pic")
        assert document.sections[0].footer._element.xpath(".//wp:anchor")
        assert not document.element.body.xpath(".//w:instrText")

        # A reference document contributes its layout, never its old body text.
        template = root / "reference.docx"
        reference = Document()
        reference.add_paragraph("OLD BODY")
        reference.sections[0].header.paragraphs[0].text = "REFERENCE HEADER"
        reference.save(template)
        subprocess.run(command + [str(template)], check=True, capture_output=True, env=env, timeout=90)
        with_template = Document(output)
        assert "OLD BODY" not in "\n".join(p.text for p in with_template.paragraphs)
        assert with_template.sections[0].header.paragraphs[0].text == "REFERENCE HEADER"
        assert "Pregunta 1. Consulta:" in "\n".join(p.text for p in with_template.paragraphs)

        entry.write_text(source.replace(r"\fontsize", r"\tableofcontents" + "\n" + r"\fontsize"), encoding="utf-8")
        subprocess.run(command, check=True, capture_output=True, env=env, timeout=90)
        indexed = Document(output)
        assert sum("TOC" in field for field in indexed.element.body.xpath(".//w:instrText/text()")) == 1
        assert "ÍNDICE DE CONTENIDOS" in "\n".join(p.text for p in indexed.paragraphs)

        # Existing headings, lists and longtable must remain usable.
        entry.write_text(r"""
\documentclass[12pt,letterpaper]{article}
\newcommand{\nombre}{\textbf{Proyecto}}
\begin{document}
\section{Título del \nombre}
Texto normal.

\begin{enumerate}[label=(\alph*)]
\item Primer elemento.
\item Segundo elemento con \nombre.
\end{enumerate}
\begin{longtable}{p{0.5\textwidth}p{0.5\textwidth}}
\toprule Nombre & Valor\\\midrule A & 10\\ B & 20\\\bottomrule
\end{longtable}
\end{document}
""", encoding="utf-8")
        subprocess.run(command, check=True, capture_output=True, env=env, timeout=90)
        regression = Document(output)
        assert any(p.text == "1  Título del Proyecto" for p in regression.paragraphs)
        assert any(p.text.startswith("(a)") and "Primer elemento" in p.text for p in regression.paragraphs)
        normal = next(p for p in regression.paragraphs if p.text == "Texto normal.")
        assert normal.runs[0].font.size.pt == 12
        assert regression.tables[0].cell(2, 1).text == "20"
        assert not regression.tables[0].cell(0, 0).paragraphs[0].paragraph_format.keep_with_next

        # Invalid input must fail before replacing an existing output.
        previous = output.read_bytes()
        entry.write_text(r"\documentclass{article}", encoding="utf-8")
        failed = subprocess.run(command, capture_output=True, env=env, timeout=90)
        assert failed.returncode != 0 and output.read_bytes() == previous
        converter.parse_preamble(r"\newcommand{\foo}{A}\renewcommand{\foo}[1]{\textbf{#1}}")
        assert converter._resolve_macros(r"\foo{B}").strip() == r"\textbf{B}"
        assert converter._extract_balanced_braces(r"{a \{b\} {c}}", 0) == r"a \{b\} {c}"
        assert converter.strip_fmt(r"\vspace{5mm} Texto") == "Texto"
        widths = converter._parse_col_widths_dxa("p{30mm}p{35mm}", 2)
        assert all(abs(width / 1440 * 25.4 - (target + pad_mm)) < 0.1 for width, target in zip(widths, (30, 35)))
        for definitions, call in ((r"\newcommand{\loop}{\loop}", r"\loop"),
                                  (r"\newcommand{\need}[2]{#1 #2}", r"\need{one}")):
            converter.parse_preamble(definitions)
            try:
                converter._resolve_macros(call)
            except ValueError:
                pass
            else:
                raise AssertionError("Invalid macro accepted")
    print("OK: content, macros, formatting, tables, equations, templates and error handling.")


if __name__ == "__main__":
    check_conversion()
